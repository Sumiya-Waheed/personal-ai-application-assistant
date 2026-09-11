"""File extraction only. No UI or RAG logic belongs here."""
from __future__ import annotations
from io import BytesIO
from pathlib import Path
import fitz
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

def extract_text(file_name: str, data: bytes) -> str:
    ext = Path(file_name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    if not data:
        raise ValueError("The uploaded file is empty.")

    if ext == ".pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            pages = [page.get_text("text") for page in doc]
        finally:
            doc.close()
        text = "\n".join(pages)
    elif ext == ".docx":
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8-sig", errors="replace")

    text = text.strip()
    if not text:
        raise ValueError(f"No extractable text was found in '{file_name}'.")
    return text
